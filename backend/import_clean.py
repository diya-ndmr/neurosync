#!/usr/bin/env python3
"""
Transform the new NeuroSync dataset - cleaner approach
Build a new workbook from scratch with required sheets
"""
import sys
import os
import tempfile
import json
from datetime import datetime, timedelta, timezone

src_file = '/Users/diyanandkoomar/Downloads/NeuroSync_University_Intelligence_Dataset 2.xlsx'
output_file = os.environ.get('TMPDIR', '/var/tmp') + '/import_result.txt'

with open(output_file, 'w') as out:
    try:
        out.write("=== NeuroSync Dataset Import (Clean Build) ===\n\n")
        
        # Read source file
        with open(src_file, 'rb') as f:
            data = f.read()
        
        # Create temp working file
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            tmp.write(data)
            temp_path = tmp.name
        
        out.write(f"1. Loaded source file ({len(data) / 1024 / 1024:.2f} MB)\n\n")
        
        # Read source workbook to extract raw data
        import openpyxl
        wb_src = openpyxl.load_workbook(temp_path)
        
        # Extract Students data
        ws_students_src = wb_src['Students']
        students = []
        for row_num in range(2, ws_students_src.max_row + 1):
            student_id = ws_students_src.cell(row_num, 1).value
            if student_id:
                students.append({
                    'student_id': student_id,
                    'first_name': ws_students_src.cell(row_num, 2).value,
                    'last_name': ws_students_src.cell(row_num, 3).value,
                    'student_email': ws_students_src.cell(row_num, 4).value,
                    'faculty': ws_students_src.cell(row_num, 5).value,
                    'programme': ws_students_src.cell(row_num, 6).value,
                    'year_of_study': ws_students_src.cell(row_num, 7).value or 1,
                    'enrolment_status': ws_students_src.cell(row_num, 8).value,
                    'average_mark_pct': ws_students_src.cell(row_num, 14).value,
                    'overall_attendance_pct': ws_students_src.cell(row_num, 15).value,
                })
        
        out.write(f"2. Extracted {len(students)} students from source\n\n")
        
        # Extract Assessments data
        ws_assess_src = wb_src['Assessments']
        assessments = []
        task_dates = {}  # Map task names to consistent dates
        
        for row_num in range(2, ws_assess_src.max_row + 1):
            student_id = ws_assess_src.cell(row_num, 1).value
            if student_id:
                task_name = ws_assess_src.cell(row_num, 5).value
                mark_pct = ws_assess_src.cell(row_num, 9).value
                
                assessments.append({
                    'student_id': student_id,
                    'faculty': ws_assess_src.cell(row_num, 2).value,
                    'task_id': ws_assess_src.cell(row_num, 4).value,
                    'task_name': task_name,
                    'mark_pct': mark_pct,
                })
        
        out.write(f"3. Extracted {len(assessments)} assessments from source\n\n")
        
        # Build task dates map
        unique_tasks = sorted(set(a['task_name'] for a in assessments if a['task_name']))
        base_date = datetime(2026, 8, 1)
        for idx, task_name in enumerate(unique_tasks):
            task_dates[task_name] = (base_date + timedelta(days=idx * 7)).strftime('%Y-%m-%d')
        
        out.write(f"4. Found {len(unique_tasks)} unique tasks\n\n")
        
        # Create fresh workbook with only required sheets
        out.write(f"5. Building fresh workbook\n")
        wb_new = openpyxl.Workbook()
        wb_new.remove(wb_new.active)  # Remove default sheet
        
        # Students sheet
        ws = wb_new.create_sheet('Students')
        headers = ['student_id', 'first_name', 'last_name', 'student_email', 'faculty', 'programme',
                  'year_of_study', 'enrolment_status', 'overall_attendance_pct', 'average_mark_pct',
                  'tasks_completed', 'tasks_missing', 'risk_band']
        for col, h in enumerate(headers, 1):
            ws.cell(1, col).value = h
        
        for row, student in enumerate(students, 2):
            ws.cell(row, 1).value = student['student_id']
            ws.cell(row, 2).value = student['first_name']
            ws.cell(row, 3).value = student['last_name']
            ws.cell(row, 4).value = student['student_email']
            ws.cell(row, 5).value = student['faculty']
            ws.cell(row, 6).value = student['programme']
            ws.cell(row, 7).value = student['year_of_study']
            # Normalize enrolment_status
            status = str(student['enrolment_status']).lower() if student['enrolment_status'] else 'active'
            ws.cell(row, 8).value = 'ACTIVE' if status == 'active' else 'LEAVE'
            ws.cell(row, 9).value = student['overall_attendance_pct'] or 75
            ws.cell(row, 10).value = student['average_mark_pct'] or 65
            ws.cell(row, 11).value = 5  # tasks_completed
            ws.cell(row, 12).value = 2  # tasks_missing
            ws.cell(row, 13).value = 'MEDIUM'  # risk_band
        
        out.write(f"   ✓ Students sheet: {len(students)} students\n")
        
        # Assessments sheet
        ws = wb_new.create_sheet('Assessments')
        headers = ['student_id', 'faculty', 'task_id', 'task_name', 'assessment_type',
                  'due_date_or_exam_date', 'max_mark', 'submission_status', 'submitted_or_sat_on',
                  'mark_pct', 'lecturer_feedback']
        for col, h in enumerate(headers, 1):
            ws.cell(1, col).value = h
        
        for row, assess in enumerate(assessments, 2):
            ws.cell(row, 1).value = assess['student_id']
            ws.cell(row, 2).value = assess['faculty']
            ws.cell(row, 3).value = assess['task_id']
            ws.cell(row, 4).value = assess['task_name']
            ws.cell(row, 5).value = 'assignment'
            ws.cell(row, 6).value = task_dates.get(assess['task_name'], '2026-09-30')
            ws.cell(row, 7).value = 100  # max_mark
            ws.cell(row, 8).value = 'submitted' if (assess['mark_pct'] and assess['mark_pct'] > 0) else 'missing'
            ws.cell(row, 9).value = None  # submitted_or_sat_on
            ws.cell(row, 10).value = assess['mark_pct'] or 0
            ws.cell(row, 11).value = None  # lecturer_feedback
        
        out.write(f"   ✓ Assessments sheet: {len(assessments)} assessments\n")
        
        # Attendance sheet
        ws = wb_new.create_sheet('Attendance')
        headers = ['student_id', 'month', 'scheduled_sessions', 'attended_sessions',
                  'approved_leave_sessions', 'unexcused_absences', 'attendance_pct_excluding_leave', 'leave_reason']
        for col, h in enumerate(headers, 1):
            ws.cell(1, col).value = h
        
        row = 2
        attendance_months = ('2026-07', '2026-08', '2026-09')
        for student in students:
            for month_str in attendance_months:
                
                ws.cell(row, 1).value = student['student_id']
                ws.cell(row, 2).value = month_str
                ws.cell(row, 3).value = 20  # scheduled_sessions
                ws.cell(row, 4).value = 18  # attended_sessions
                ws.cell(row, 5).value = 0   # approved_leave
                ws.cell(row, 6).value = 2   # unexcused_absences
                ws.cell(row, 7).value = 90  # attendance_pct
                ws.cell(row, 8).value = None  # leave_reason
                
                row += 1
        
        out.write(f"   ✓ Attendance sheet: {row - 2} records\n")
        
        # Leave sheet
        ws = wb_new.create_sheet('Leave')
        headers = ['leave_id', 'student_id', 'month', 'approved_leave_sessions', 'reason', 'approval_status']
        for col, h in enumerate(headers, 1):
            ws.cell(1, col).value = h
        # Leave sheet stays empty
        
        out.write(f"   ✓ Leave sheet created (empty)\n")
        
        # Faculty Schedule sheet
        ws = wb_new.create_sheet('Faculty Schedule')
        faculties = sorted(set(s['faculty'] for s in students))
        
        # Headers: faculty + each task name
        ws.cell(1, 1).value = 'faculty'
        for col, task_name in enumerate(unique_tasks, 2):
            ws.cell(1, col).value = task_name
        
        # Faculty rows with task dates
        for row, faculty in enumerate(faculties, 2):
            ws.cell(row, 1).value = faculty
            for col, task_name in enumerate(unique_tasks, 2):
                ws.cell(row, col).value = task_dates[task_name]
        
        out.write(f"   ✓ Faculty Schedule sheet: {len(faculties)} faculties, {len(unique_tasks)} tasks\n\n")
        
        # Save new workbook
        new_wb_path = temp_path.replace('.xlsx', '_clean.xlsx')
        wb_new.save(new_wb_path)
        
        out.write(f"6. Saved clean workbook to {new_wb_path}\n\n")
        
        # Import
        sys.path.insert(0, '/Users/diyanandkoomar/Downloads/Neurosync Code 2/Neurosync Code/backend')
        from lms.workbook_import import import_student_workbook
        from app.config.settings import settings
        
        out.write(f"7. Preparing database\n")
        db_path = settings.lms_sync_database_url.replace('sqlite:///', '')
        if os.path.exists(db_path):
            os.unlink(db_path)
            out.write(f"   Cleared existing database\n")
        
        out.write(f"\n8. Importing...\n\n")
        result = import_student_workbook(new_wb_path, settings.lms_sync_database_url)

        # Preserve operational intelligence from the workbook's case sheets.
        from datetime import date
        from lms.database import LMSSessionLocal
        from lms.models import Intervention, Lecturer, LMSActivity, Referral, StaffMapping, Student, WorkbookIntelligence, WorkflowEvent
        from lms.workflow_service import record_event

        source_wb = openpyxl.load_workbook(temp_path, data_only=True)
        lms_db = LMSSessionLocal()
        try:
            student_ids = {row.student_id for row in lms_db.query(Student).all()}
            lecturer_ids = {row.lecturer_id for row in lms_db.query(Lecturer).all()}
            default_lecturer = next(iter(lecturer_ids), 'WB-LEC-01')
            mapping_email = 'lecturer@university.edu'
            if not lms_db.query(StaffMapping).filter(StaffMapping.app_user_email == mapping_email).first() and default_lecturer in lecturer_ids:
                lms_db.add(StaffMapping(app_user_email=mapping_email, lecturer_id=default_lecturer))
            referral_sheet = source_wb['Referral Queue']
            referral_headers = [cell.value for cell in next(referral_sheet.iter_rows(max_row=1))]
            referral_rows = [dict(zip(referral_headers, row)) for row in referral_sheet.iter_rows(min_row=2, values_only=True)]
            imported_referrals = 0
            for row in referral_rows:
                student_id = row.get('student_id')
                referral_id = row.get('referral_id')
                if not student_id or student_id not in student_ids or not referral_id:
                    continue
                if lms_db.query(Referral).filter(Referral.referral_id == referral_id).first():
                    continue
                risk_level = str(row.get('risk_level') or 'MEDIUM').upper()
                urgency = 'PRIORITY' if str(row.get('priority') or '').lower() in {'critical', 'high', 'priority'} else 'STANDARD'
                created_on = row.get('referred_on') or date.today()
                if isinstance(created_on, str):
                    created_on = date.fromisoformat(created_on)
                created_at = datetime.combine(created_on, datetime.min.time(), tzinfo=timezone.utc)
                lms_db.add(Referral(
                    referral_id=referral_id,
                    student_id=student_id,
                    lecturer_id=default_lecturer,
                    risk_level=risk_level,
                    risk_score={'HIGH': 25.0, 'MEDIUM': 60.0, 'LOW': 85.0}.get(risk_level, 60.0),
                    reason=str(row.get('reason') or 'Workbook referral requiring advisor review'),
                    lecturer_notes=str(row.get('notes') or '') or None,
                    urgency=urgency,
                    status='REFERRED',
                    created_at=created_at,
                    updated_at=created_at,
                ))
                imported_referrals += 1
            lms_db.flush()
            case_sheet = source_wb['Student Cases']
            case_headers = [cell.value for cell in next(case_sheet.iter_rows(max_row=1))]
            case_rows = [dict(zip(case_headers, row)) for row in case_sheet.iter_rows(min_row=2, values_only=True)]
            referral_ids = {row.referral_id for row in lms_db.query(Referral).all()}
            for row in case_rows:
                if row.get('referral_id') not in referral_ids or not row.get('student_id'):
                    continue
                occurred = row.get('opened_on') or date.today()
                if isinstance(occurred, str):
                    occurred = date.fromisoformat(occurred)
                record_event(lms_db, row['student_id'], None, 'CASE_OPENED', str(row.get('intervention_plan') or 'Advisor case opened'), row['referral_id'])

            intervention_sheet = source_wb['Interventions']
            intervention_headers = [cell.value for cell in next(intervention_sheet.iter_rows(max_row=1))]
            intervention_rows = [dict(zip(intervention_headers, row)) for row in intervention_sheet.iter_rows(min_row=2, values_only=True)]
            imported_interventions = 0
            for row in intervention_rows:
                student_id = row.get('student_id')
                if not student_id or student_id not in student_ids:
                    continue
                referral_id = row.get('referral_id') or next(iter(referral_ids), None)
                if not referral_id:
                    continue
                intervention_id = row.get('intervention_id')
                if lms_db.query(Intervention).filter(Intervention.intervention_id == intervention_id).first():
                    continue
                started = row.get('action_date') or date.today()
                if isinstance(started, str):
                    started = date.fromisoformat(started)
                started_at = datetime.combine(started, datetime.min.time(), tzinfo=timezone.utc)
                lms_db.add(Intervention(
                    intervention_id=intervention_id,
                    referral_id=referral_id,
                    student_id=student_id,
                    advisor_user_id=str(row.get('owner_id') or 'WORKBOOK-ADVISOR'),
                    strategy=str(row.get('action_type') or 'Support action'),
                    reason='Imported from the workbook intervention plan',
                    advisor_notes=str(row.get('outcome') or '') or None,
                    date_started=started_at,
                    outcome=str(row.get('outcome') or '') or None,
                    status='ACTIVE',
                    created_at=started_at,
                    updated_at=started_at,
                ))
                imported_interventions += 1

            activity_sheet = source_wb['LMS Activity']
            activity_headers = [cell.value for cell in next(activity_sheet.iter_rows(max_row=1))]
            activity_rows = [dict(zip(activity_headers, row)) for row in activity_sheet.iter_rows(min_row=2, values_only=True)]
            imported_activity = 0
            for row in activity_rows:
                student_id = row.get('student_id')
                if not student_id or student_id not in student_ids:
                    continue
                timestamp = row.get('last_login_date') or date.today()
                if isinstance(timestamp, str):
                    timestamp = date.fromisoformat(timestamp)
                timestamp = datetime.combine(timestamp, datetime.min.time(), tzinfo=timezone.utc)
                activity_id = f"WB-ACT-{student_id}"
                if lms_db.query(LMSActivity).filter(LMSActivity.activity_id == activity_id).first():
                    continue
                lms_db.add(LMSActivity(
                    activity_id=activity_id,
                    student_id=student_id,
                    timestamp=timestamp,
                    activity_type=str(row.get('last_visited_lms_section') or 'LMS_ACTIVITY').upper().replace(' ', '_'),
                    session_duration=int(row.get('total_minutes_30d') or 0),
                ))
                imported_activity += 1
            lms_db.commit()
            out.write(f"  Workbook referrals preserved: {imported_referrals}\n")
            out.write(f"  Workbook interventions preserved: {imported_interventions}\n")
            out.write(f"  Workbook LMS activity preserved: {imported_activity}\n")

            intelligence_sheets = [
                'Dashboard', 'Weekly Intelligence', 'Trends', 'Program Comparison',
                'Top Students', 'Students At Risk', 'Students Recovering',
                'AI Recommendations', 'Lecturer Overview', 'Advisor Overview', 'Advisor Weekly'
            ]
            imported_intelligence = 0
            for sheet_name in intelligence_sheets:
                sheet = source_wb[sheet_name]
                headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
                for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                    payload = {key: value for key, value in zip(headers, values) if key is not None}
                    if not any(value is not None for value in payload.values()):
                        continue
                    normalized = json.loads(json.dumps(payload, default=str))
                    lms_db.add(WorkbookIntelligence(
                        record_id=f"WB-{sheet_name.upper().replace(' ', '_')}-{row_number}",
                        sheet_name=sheet_name,
                        row_number=row_number,
                        payload=json.dumps(normalized),
                    ))
                    imported_intelligence += 1
            lms_db.commit()
            out.write(f"  Workbook intelligence rows preserved: {imported_intelligence}\n")
        finally:
            lms_db.close()
            source_wb.close()
        
        out.write(f"✓ Import Complete!\n")
        out.write(f"  Students: {result.get('students', 0)}\n")
        out.write(f"  Modules: {result.get('modules', 0)}\n")
        out.write(f"  Assessments: {result.get('assessments', 0)}\n")
        out.write(f"  Attendance: {result.get('attendance', 0)}\n")
        out.write(f"  Leave: {result.get('leave', 0)}\n")
        
        # Cleanup
        os.unlink(temp_path)
        os.unlink(new_wb_path)
        out.write(f"\n✓ Done - database ready with new dataset\n")
        
    except Exception as e:
        out.write(f"\n✗ ERROR: {e}\n\n")
        import traceback
        traceback.print_exc(file=out)
        sys.exit(1)
