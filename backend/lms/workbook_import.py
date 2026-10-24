from __future__ import annotations

import hashlib
import re
import uuid
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from lms.database import LMSBase, build_lms_session_factory, ensure_lms_db
from lms.models import (
    ApprovedLeaveRecord,
    Assessment,
    AssessmentResult,
    DatasetImport,
    Lecturer,
    Module,
    ModuleEnrolment,
    MonthlyAttendanceSummary,
    Programme,
    StaffMapping,
    Student,
    StudentMapping,
    Submission,
)

REQUIRED_HEADERS = {
    "Students": {
        "student_id", "first_name", "last_name", "student_email", "faculty", "programme",
        "year_of_study", "enrolment_status", "overall_attendance_pct", "average_mark_pct",
        "tasks_completed", "tasks_missing", "risk_band",
    },
    "Attendance": {
        "student_id", "month", "scheduled_sessions", "attended_sessions", "approved_leave_sessions",
        "unexcused_absences", "attendance_pct_excluding_leave", "leave_reason",
    },
    "Assessments": {
        "student_id", "faculty", "task_id", "task_name", "assessment_type", "due_date_or_exam_date",
        "max_mark", "submission_status", "submitted_or_sat_on", "mark_pct", "lecturer_feedback",
    },
    "Leave": {"leave_id", "student_id", "month", "approved_leave_sessions", "reason", "approval_status"},
    "Faculty Schedule": {"faculty"},
}

TASK_STATUSES = {
    "submitted": "SUBMITTED",
    "submitted late": "LATE",
    "sat exam": "SUBMITTED",
    "missing": "MISSING",
    "missed exam": "MISSING",
}


def _read_sheet(workbook, sheet_name: str) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Workbook is missing required sheet: {sheet_name}")
    rows = workbook[sheet_name].iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        raise ValueError(f"Sheet {sheet_name} is empty")
    normalized_headers = [str(value).strip() if value is not None else "" for value in headers]
    missing = REQUIRED_HEADERS[sheet_name] - set(normalized_headers)
    if missing:
        raise ValueError(f"Sheet {sheet_name} is missing columns: {', '.join(sorted(missing))}")
    result = []
    for row in rows:
        if not any(value is not None for value in row):
            continue
        result.append(dict(zip(normalized_headers, row)))
    return result


def _text(value: Any, field: str) -> str:
    if value is None or not str(value).strip():
        raise ValueError(f"Required workbook value is missing: {field}")
    return str(value).strip()


def _integer(value: Any, field: str, minimum: int = 0) -> int:
    try:
        number = int(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"Invalid integer value in {field}: {value!r}") from error
    if number < minimum:
        raise ValueError(f"Value in {field} must be at least {minimum}")
    return number


def _number(value: Any, field: str, minimum: float, maximum: float) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"Invalid numeric value in {field}: {value!r}") from error
    if not minimum <= number <= maximum:
        raise ValueError(f"Value in {field} must be between {minimum} and {maximum}")
    return number


def _date(value: Any, field: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except (TypeError, ValueError) as error:
        raise ValueError(f"Invalid date in {field}: {value!r}") from error


def _slug(value: str) -> str:
    slug = re.sub(r"[^A-Z0-9]+", "-", value.upper()).strip("-")
    return slug[:40] or "UNKNOWN"


def _load_and_validate(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {name: _read_sheet(workbook, name) for name in REQUIRED_HEADERS}
    finally:
        workbook.close()

    students = sheets["Students"]
    attendance = sheets["Attendance"]
    assessments = sheets["Assessments"]
    leave = sheets["Leave"]
    schedules = sheets["Faculty Schedule"]

    student_ids = [_text(row["student_id"], "Students.student_id") for row in students]
    if len(student_ids) != len(set(student_ids)):
        raise ValueError("Students contains duplicate student_id values")
    student_id_set = set(student_ids)
    email_values = [_text(row["student_email"], "Students.student_email").lower() for row in students]
    if len(email_values) != len(set(email_values)):
        raise ValueError("Students contains duplicate student_email values")

    student_faculties = {}
    student_programmes = {}
    programme_faculties = {}
    for row in students:
        student_id = _text(row["student_id"], "Students.student_id")
        faculty = _text(row["faculty"], "Students.faculty")
        programme = _text(row["programme"], "Students.programme")
        student_faculties[student_id] = faculty
        student_programmes[student_id] = programme
        if programme in programme_faculties and programme_faculties[programme] != faculty:
            raise ValueError(f"Programme {programme!r} is associated with multiple faculties")
        programme_faculties[programme] = faculty
        _number(row["overall_attendance_pct"], "Students.overall_attendance_pct", 0, 100)
        if row["average_mark_pct"] is not None:
            _number(row["average_mark_pct"], "Students.average_mark_pct", 0, 100)
        _integer(row["year_of_study"], "Students.year_of_study", 1)
        if _text(row["enrolment_status"], "Students.enrolment_status").casefold() not in {"active", "leave", "suspended", "graduated"}:
            raise ValueError(f"Unsupported enrolment status for {student_id}")

    faculty_schedules = {}
    for row in schedules:
        faculty = _text(row["faculty"], "Faculty Schedule.faculty")
        if faculty in faculty_schedules:
            raise ValueError(f"Duplicate faculty schedule for {faculty!r}")
        faculty_schedules[faculty] = {
            key: _date(value, f"Faculty Schedule.{faculty}.{key}")
            for key, value in row.items() if key != "faculty" and value is not None
        }
    if set(student_faculties.values()) - set(faculty_schedules):
        raise ValueError("Every student faculty must have a Faculty Schedule row")

    task_definitions = {}
    for row in assessments:
        student_id = _text(row["student_id"], "Assessments.student_id")
        if student_id not in student_id_set:
            raise ValueError(f"Assessment references unknown student {student_id}")
        faculty = _text(row["faculty"], "Assessments.faculty")
        if faculty != student_faculties[student_id]:
            raise ValueError(f"Assessment faculty does not match student {student_id}")
        task_id = _text(row["task_id"], "Assessments.task_id")
        task_name = _text(row["task_name"], "Assessments.task_name")
        assessment_type = _text(row["assessment_type"], "Assessments.assessment_type")
        status_label = _text(row["submission_status"], "Assessments.submission_status").casefold()
        if status_label not in TASK_STATUSES:
            raise ValueError(f"Unsupported assessment status {status_label!r}")
        task_key = (faculty, task_id)
        task_signature = (task_name, assessment_type)
        if task_key in task_definitions and task_definitions[task_key] != task_signature:
            raise ValueError(f"Conflicting task definitions for {task_key}")
        task_definitions[task_key] = task_signature
        due_date = _date(row["due_date_or_exam_date"], "Assessments.due_date_or_exam_date")
        schedule_date = faculty_schedules[faculty].get(task_name)
        if schedule_date is not None and schedule_date != due_date:
            raise ValueError(f"Assessment date conflicts with faculty schedule for {faculty} / {task_name}")
        _number(row["max_mark"], "Assessments.max_mark", 0.01, 100000)
        if row["mark_pct"] is not None:
            _number(row["mark_pct"], "Assessments.mark_pct", 0, 100)
        if row["submitted_or_sat_on"] is not None:
            _date(row["submitted_or_sat_on"], "Assessments.submitted_or_sat_on")

    for row in attendance:
        student_id = _text(row["student_id"], "Attendance.student_id")
        if student_id not in student_id_set:
            raise ValueError(f"Attendance references unknown student {student_id}")
        month = _text(row["month"], "Attendance.month")
        if not re.fullmatch(r"\d{4}-\d{2}", month):
            raise ValueError(f"Attendance month must use YYYY-MM: {month!r}")
        scheduled = _integer(row["scheduled_sessions"], "Attendance.scheduled_sessions")
        attended = _integer(row["attended_sessions"], "Attendance.attended_sessions")
        leave_sessions = _integer(row["approved_leave_sessions"], "Attendance.approved_leave_sessions")
        absent = _integer(row["unexcused_absences"], "Attendance.unexcused_absences")
        if attended + leave_sessions + absent > scheduled:
            raise ValueError(f"Attendance counts exceed scheduled sessions for {student_id} in {month}")
        _number(row["attendance_pct_excluding_leave"], "Attendance.attendance_pct_excluding_leave", 0, 100)

    leave_ids = set()
    for row in leave:
        leave_id = _text(row["leave_id"], "Leave.leave_id")
        if leave_id in leave_ids:
            raise ValueError(f"Duplicate leave_id {leave_id}")
        leave_ids.add(leave_id)
        if _text(row["student_id"], "Leave.student_id") not in student_id_set:
            raise ValueError(f"Leave references unknown student {row['student_id']}")
        approval = _text(row["approval_status"], "Leave.approval_status").casefold()
        if approval != "approved":
            raise ValueError(f"Workbook includes non-approved leave record {leave_id}")
        _integer(row["approved_leave_sessions"], "Leave.approved_leave_sessions", 1)

    faculty_indexes = {faculty: index for index, faculty in enumerate(sorted(faculty_schedules), start=1)}
    task_indexes = {task_id: index for index, task_id in enumerate(sorted({task_id for _, task_id in task_definitions}), start=1)}
    if max(task_indexes.values(), default=0) > 99 or max(faculty_indexes.values(), default=0) > 99:
        raise ValueError("Workbook has too many faculties or tasks for canonical module IDs")

    return {
        "students": students,
        "attendance": attendance,
        "assessments": assessments,
        "leave": leave,
        "schedules": faculty_schedules,
        "programme_faculties": programme_faculties,
        "student_faculties": student_faculties,
        "student_programmes": student_programmes,
        "task_definitions": task_definitions,
        "faculty_indexes": faculty_indexes,
        "task_indexes": task_indexes,
    }


def import_student_workbook(workbook_path: str, database_url: str) -> dict[str, Any]:
    source = Path(workbook_path).expanduser().resolve(strict=True)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    data = _load_and_validate(source)
    target_url = database_url
    engine, session_factory = build_lms_session_factory(database_url=target_url)
    try:
        LMSBase.metadata.create_all(bind=engine)
        db: Session = session_factory()
        try:
            previous = db.query(DatasetImport).filter(DatasetImport.content_sha256 == digest).first()
            if previous:
                if db.query(Student).count() != previous.student_count:
                    raise ValueError("Import target changed after this dataset import; refusing an ambiguous mixed database")
                demo_student_mapping = db.query(StudentMapping).filter(
                    StudentMapping.app_user_email == "sarah@university.edu"
                ).first()
                if demo_student_mapping is not None:
                    db.delete(demo_student_mapping)
                demo_staff_mapping = db.query(StaffMapping).filter(
                    StaffMapping.app_user_email == "lecturer@university.edu"
                ).first()
                if demo_staff_mapping is not None:
                    db.delete(demo_staff_mapping)
                db.commit()
                return {
                    "status": "already_imported",
                    "import_id": previous.import_id,
                    "student_count": previous.student_count,
                    "attendance_count": previous.attendance_count,
                    "assessment_count": previous.assessment_count,
                    "leave_count": previous.leave_count,
                }
            if db.query(Student).count() > 0:
                raise ValueError("Import target already contains students; use a new empty LMS database to avoid merging datasets")

            faculty_to_lecturer = {}
            for faculty, index in data["faculty_indexes"].items():
                lecturer_id = f"WB-LEC-{index:02d}"
                faculty_to_lecturer[faculty] = lecturer_id
                db.add(Lecturer(
                    lecturer_id=lecturer_id,
                    first_name="Dataset",
                    surname=f"Faculty {index}",
                    email=f"workbook-faculty-{index:02d}@neurosync.invalid",
                    department=faculty,
                ))

            programme_ids = {}
            for programme, faculty in sorted(data["programme_faculties"].items()):
                programme_id = f"WB-PRG-{_slug(programme)}"
                programme_ids[programme] = programme_id
                db.add(Programme(
                    programme_id=programme_id,
                    programme_code=f"WB-{_slug(programme)}",
                    programme_name=programme,
                    faculty=faculty,
                    level="Undergraduate",
                ))

            programme_indexes = {programme: index for index, programme in enumerate(sorted(programme_ids), start=1)}
            module_ids = {}
            for programme, faculty in sorted(data["programme_faculties"].items()):
                for (task_faculty, task_id), task_definition in sorted(data["task_definitions"].items()):
                    if task_faculty != faculty:
                        continue
                    task_name = task_definition[0]
                    module_id = f"MOD{programme_indexes[programme]:02d}{data['task_indexes'][task_id]:02d}"
                    module_ids[(programme, task_id)] = module_id
                    db.add(Module(
                        module_id=module_id,
                        programme_id=programme_ids[programme],
                        module_code=f"WB{programme_indexes[programme]:02d}{data['task_indexes'][task_id]:02d}",
                        module_name=task_name,
                        credits=20,
                        semester=1,
                        lecturer_id=faculty_to_lecturer[faculty],
                    ))
            db.flush()

            for row in data["students"]:
                student_id = _text(row["student_id"], "Students.student_id")
                programme = _text(row["programme"], "Students.programme")
                band = _text(row["risk_band"], "Students.risk_band").casefold()
                status = _text(row["enrolment_status"], "Students.enrolment_status").upper()
                student = Student(
                    student_id=student_id,
                    first_name=_text(row["first_name"], "Students.first_name"),
                    surname=_text(row["last_name"], "Students.last_name"),
                    institutional_email=_text(row["student_email"], "Students.student_email").lower(),
                    programme_id=programme_ids[programme],
                    year_of_study=_integer(row["year_of_study"], "Students.year_of_study", 1),
                    semester=1,
                    behaviour_group=f"imported_{_slug(band).lower()}",
                    enrolment_status=status,
                )
                db.add(student)
                for module_key, module_id in module_ids.items():
                    if module_key[0] == programme:
                        db.add(ModuleEnrolment(student_id=student_id, module_id=module_id, status="ENROLLED"))

            schedule_rows = data["schedules"]
            assessment_ids = {}
            for programme, faculty in data["programme_faculties"].items():
                for (task_faculty, task_id), (task_name, assessment_type) in data["task_definitions"].items():
                    if task_faculty != faculty:
                        continue
                    module_id = module_ids[(programme, task_id)]
                    due_date = _date(schedule_rows[faculty][task_name], f"Faculty Schedule.{faculty}.{task_name}")
                    assessment_id = f"WB-ASM-{programme_indexes[programme]:02d}-{data['task_indexes'][task_id]:02d}"
                    assessment_ids[(programme, task_id)] = assessment_id
                    db.add(Assessment(
                        assessment_id=assessment_id,
                        module_id=module_id,
                        title=task_name,
                        assessment_type=_slug(assessment_type),
                        weight=20.0,
                        maximum_mark=100.0,
                        due_date=due_date,
                    ))

            for row in data["assessments"]:
                student_id = _text(row["student_id"], "Assessments.student_id")
                faculty = _text(row["faculty"], "Assessments.faculty")
                task_id = _text(row["task_id"], "Assessments.task_id")
                programme = next(student["programme"] for student in data["students"] if student["student_id"] == student_id)
                assessment_id = assessment_ids[(programme, task_id)]
                status = TASK_STATUSES[_text(row["submission_status"], "Assessments.submission_status").casefold()]
                submitted_date = _date(row["submitted_or_sat_on"], "Assessments.submitted_or_sat_on") if row["submitted_or_sat_on"] is not None else None
                submission_id = f"WB-SUB-{_slug(student_id)}-{data['task_indexes'][task_id]:02d}"
                db.add(Submission(
                    submission_id=submission_id,
                    student_id=student_id,
                    assessment_id=assessment_id,
                    submitted_at=datetime.combine(submitted_date, datetime.min.time()) if submitted_date else None,
                    status=status,
                ))
                if row["mark_pct"] is not None:
                    percentage = _number(row["mark_pct"], "Assessments.mark_pct", 0, 100)
                    maximum_mark = _number(row["max_mark"], "Assessments.max_mark", 0.01, 100000)
                    db.add(AssessmentResult(
                        result_id=f"WB-RES-{_slug(student_id)}-{data['task_indexes'][task_id]:02d}",
                        student_id=student_id,
                        assessment_id=assessment_id,
                        mark=round(maximum_mark * percentage / 100, 2),
                        percentage=percentage,
                        submission_date=submitted_date,
                        status=status,
                    ))

            for row in data["attendance"]:
                student_id = _text(row["student_id"], "Attendance.student_id")
                month = _text(row["month"], "Attendance.month")
                db.add(MonthlyAttendanceSummary(
                    summary_id=f"WB-ATT-{_slug(student_id)}-{month.replace('-', '')}",
                    student_id=student_id,
                    month=month,
                    scheduled_sessions=_integer(row["scheduled_sessions"], "Attendance.scheduled_sessions"),
                    attended_sessions=_integer(row["attended_sessions"], "Attendance.attended_sessions"),
                    approved_leave_sessions=_integer(row["approved_leave_sessions"], "Attendance.approved_leave_sessions"),
                    unexcused_absences=_integer(row["unexcused_absences"], "Attendance.unexcused_absences"),
                    attendance_percentage=_number(row["attendance_pct_excluding_leave"], "Attendance.attendance_pct_excluding_leave", 0, 100),
                    source="NeuroSync_500_Student_Dataset",
                ))
            for row in data["leave"]:
                db.add(ApprovedLeaveRecord(
                    leave_id=_text(row["leave_id"], "Leave.leave_id"),
                    student_id=_text(row["student_id"], "Leave.student_id"),
                    month=_text(row["month"], "Leave.month"),
                    sessions=_integer(row["approved_leave_sessions"], "Leave.approved_leave_sessions", 1),
                    reason=str(row["reason"]).strip() if row["reason"] is not None else None,
                    approval_status=_text(row["approval_status"], "Leave.approval_status").upper(),
                    source="NeuroSync_500_Student_Dataset",
                ))

            manifest = DatasetImport(
                import_id=f"IMP-{uuid.uuid4().hex[:12].upper()}",
                source_name=source.name,
                content_sha256=digest,
                student_count=len(data["students"]),
                attendance_count=len(data["attendance"]),
                assessment_count=len(data["assessments"]),
                leave_count=len(data["leave"]),
            )
            db.add(manifest)
            db.commit()
            return {
                "status": "imported",
                "import_id": manifest.import_id,
                "student_count": manifest.student_count,
                "attendance_count": manifest.attendance_count,
                "assessment_count": manifest.assessment_count,
                "leave_count": manifest.leave_count,
                "module_count": len(module_ids),
            }
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()
    finally:
        engine.dispose()


def main() -> None:
    import argparse

    from lms.database import CONFIGURED_LMS_DATABASE_URL

    parser = argparse.ArgumentParser(description="Import NeuroSync's canonical 500-student Excel dataset.")
    parser.add_argument("workbook", help="Path to NeuroSync_500_Student_Dataset.xlsx")
    parser.add_argument(
        "--database-url",
        default=CONFIGURED_LMS_DATABASE_URL,
        help="Dedicated empty LMS database URL; defaults to MOCK_LMS_DATABASE_URL.",
    )
    parser.add_argument(
        "--use-existing-lms-db",
        action="store_true",
        help="Explicitly target the default mock LMS SQLite file; only works if it contains no students.",
    )
    args = parser.parse_args()
    database_url = args.database_url
    if args.use_existing_lms_db:
        database_url = f"sqlite:///{ensure_lms_db()}"
    if not database_url:
        parser.error("Set MOCK_LMS_DATABASE_URL or pass --database-url for a dedicated import database")
    if database_url == f"sqlite:///{ensure_lms_db()}" and not args.use_existing_lms_db:
        parser.error("Refusing to import into the synthetic default LMS database without --use-existing-lms-db")
    result = import_student_workbook(args.workbook, database_url)
    print(result)


if __name__ == "__main__":
    main()
