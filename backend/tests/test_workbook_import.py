import pytest
from datetime import date
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import lms.database as lms_database
from lms.database import LMSBase
from lms.generator import lms_dataset_is_seeded
from lms.models import (
    ApprovedLeaveRecord,
    Assessment,
    AssessmentResult,
    DatasetImport,
    LMSActivity,
    Module,
    MonthlyAttendanceSummary,
    StaffMapping,
    Student,
    StudentMapping,
    Submission,
)
from lms.workbook_import import import_student_workbook
from lms.workflow_service import build_student_review, ensure_fictional_student_emails
from app.schemas.workflow_schema import ReferralCreate

TASKS = [
    ("TASK-1", "Quiz 1", "Quiz"),
    ("TASK-2", "Assignment 1", "Assignment"),
    ("TASK-3", "Practical", "Practical"),
    ("TASK-4", "Assignment 2", "Assignment"),
    ("TASK-5", "Final Exam", "Exam"),
]


def _write_workbook(path):
    workbook = Workbook()
    students_sheet = workbook.active
    assert students_sheet is not None
    students_sheet.title = "Students"
    students_sheet.append([
        "student_id", "first_name", "last_name", "student_email", "faculty", "programme",
        "year_of_study", "enrolment_status", "advisor_id", "overall_attendance_pct",
        "average_mark_pct", "tasks_completed", "tasks_missing", "risk_band",
    ])
    students_sheet.append(["STU-2026-0001", "Avery", "Example", "avery@example.edu", "Faculty A", "Programme One", 1, "Active", "ADV-1", 88, 75, 4, 1, "Moderate"])
    students_sheet.append(["STU-2026-0002", "Jordan", "Example", "jordan@example.edu", "Faculty A", "Programme Two", 2, "Active", "ADV-2", 91, None, 5, 0, "Low"])

    attendance_sheet = workbook.create_sheet("Attendance")
    attendance_sheet.append(["student_id", "month", "scheduled_sessions", "attended_sessions", "approved_leave_sessions", "unexcused_absences", "attendance_pct_excluding_leave", "leave_reason"])
    attendance_sheet.append(["STU-2026-0001", "2026-08", 10, 8, 1, 1, 88.9, None])
    attendance_sheet.append(["STU-2026-0002", "2026-08", 10, 9, 0, 1, 90, None])

    assessments_sheet = workbook.create_sheet("Assessments")
    assessments_sheet.append(["student_id", "faculty", "task_id", "task_name", "assessment_type", "due_date_or_exam_date", "max_mark", "submission_status", "submitted_or_sat_on", "mark_pct", "lecturer_feedback"])
    for student_id in ("STU-2026-0001", "STU-2026-0002"):
        for index, (task_id, title, assessment_type) in enumerate(TASKS):
            status = "Missing" if student_id.endswith("0001") and task_id == "TASK-5" else "Submitted"
            submitted = None if status == "Missing" else date(2026, 8, min(index + 1, 28))
            mark = None if status == "Missing" else 70 + index
            assessments_sheet.append([student_id, "Faculty A", task_id, title, assessment_type, date(2026, 8, min(index * 2 + 2, 28)), 100, status, submitted, mark, None])

    leave_sheet = workbook.create_sheet("Leave")
    leave_sheet.append(["leave_id", "student_id", "month", "approved_leave_sessions", "reason", "approval_status"])
    leave_sheet.append(["LEA-TEST-1", "STU-2026-0001", "2026-08", 1, "Approved leave", "Approved"])

    schedule_sheet = workbook.create_sheet("Faculty Schedule")
    schedule_sheet.append(["faculty", "Quiz 1", "Assignment 1", "Practical", "Assignment 2", "Final Exam"])
    schedule_sheet.append(["Faculty A", date(2026, 8, 2), date(2026, 8, 4), date(2026, 8, 6), date(2026, 8, 8), date(2026, 8, 10)])

    readme = workbook.create_sheet("Read Me")
    readme.append(["Dataset", "Test fixture"])
    workbook.save(path)
    return path


@pytest.fixture

def imported_workbook(tmp_path):
    workbook_path = _write_workbook(tmp_path / "dataset.xlsx")
    database_url = f"sqlite:///{tmp_path / 'workbook.db'}"
    summary = import_student_workbook(str(workbook_path), database_url)
    engine = create_engine(database_url)
    factory = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    db = factory()
    yield summary, database_url, db
    db.close()
    engine.dispose()


def test_workbook_import_maps_all_sheets_without_fabricating_lms_activity(imported_workbook):
    summary, database_url, db = imported_workbook

    assert summary["status"] == "imported"
    assert summary["student_count"] == 2
    assert summary["attendance_count"] == 2
    assert summary["assessment_count"] == 10
    assert summary["leave_count"] == 1
    assert db.query(Student).count() == 2
    assert db.query(Module).count() == 10
    assert db.query(Assessment).count() == 10
    assert db.query(Submission).count() == 10
    assert db.query(AssessmentResult).count() == 9
    assert db.query(MonthlyAttendanceSummary).count() == 2
    assert db.query(ApprovedLeaveRecord).count() == 1
    assert db.query(LMSActivity).count() == 0
    assert summary["module_count"] == 10

    student = db.query(Student).filter(Student.student_id == "STU-2026-0001").one()
    assert student.institutional_email == "avery@example.edu"
    assert student.programme.programme_name == "Programme One"
    assert db.query(AssessmentResult).filter(AssessmentResult.student_id == student.student_id, AssessmentResult.status == "MISSING").count() == 0
    assert db.query(Submission).filter(Submission.student_id == student.student_id, Submission.status == "MISSING").count() == 1
    assert db.query(AssessmentResult).filter(AssessmentResult.student_id == "STU-2026-0002").count() == 5
    assert ReferralCreate(student_id="STU-2026-0001", reason="Review recorded attendance and course results.")
    ensure_fictional_student_emails(db)
    assert db.query(StudentMapping).filter(StudentMapping.app_user_email == "sarah@university.edu").count() == 0
    assert db.query(StaffMapping).filter(StaffMapping.app_user_email == "lecturer@university.edu").count() == 0
    review = build_student_review(db, student.student_id)
    assert len(review["attendance"]) == 1
    assert all(row["session"] == "MONTHLY SUMMARY" for row in review["attendance"])
    assert db.query(DatasetImport).count() == 1
    original_engine = lms_database.lms_engine
    original_factory = lms_database.LMSSessionLocal
    try:
        lms_database.lms_engine, lms_database.LMSSessionLocal = lms_database.build_lms_session_factory(database_url=database_url)
        assert lms_dataset_is_seeded()
    finally:
        lms_database.lms_engine.dispose()
        lms_database.lms_engine = original_engine
        lms_database.LMSSessionLocal = original_factory


def test_workbook_import_is_idempotent_and_refuses_mixed_nonempty_database(imported_workbook, tmp_path):
    summary, database_url, db = imported_workbook
    workbook_path = tmp_path / "dataset.xlsx"
    db.add(StudentMapping(app_user_email="sarah@university.edu", student_id="STU-2026-0001"))
    db.add(StaffMapping(
        app_user_email="lecturer@university.edu",
        lecturer_id=db.query(Module).first().lecturer_id,
    ))
    db.commit()
    db.close()

    repeated = import_student_workbook(str(workbook_path), database_url)
    assert repeated["status"] == "already_imported"
    assert repeated["import_id"] == summary["import_id"]

    db = sessionmaker(bind=create_engine(database_url))()
    assert db.query(StudentMapping).filter(StudentMapping.app_user_email == "sarah@university.edu").count() == 0
    assert db.query(StaffMapping).filter(StaffMapping.app_user_email == "lecturer@university.edu").count() == 0
    db.close()
    db = sessionmaker(bind=create_engine(database_url))()
    db.add(Student(
        student_id="UNRELATED", first_name="Other", surname="Record",
        institutional_email="other@example.edu", programme_id="unmapped",
        year_of_study=1, semester=1,
    ))
    db.commit()
    db.close()
    with pytest.raises(ValueError, match="mixed database"):
        import_student_workbook(str(workbook_path), database_url)


def test_workbook_validation_rejects_nonempty_target_without_partial_writes(tmp_path):
    workbook_path = _write_workbook(tmp_path / "dataset.xlsx")
    database_url = f"sqlite:///{tmp_path / 'nonempty.db'}"
    engine = create_engine(database_url)
    LMSBase.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    db.add(Student(
        student_id="EXISTING", first_name="Existing", surname="Student",
        institutional_email="existing@example.edu", programme_id="unmapped",
        year_of_study=1, semester=1,
    ))
    db.commit()
    db.close()
    engine.dispose()

    with pytest.raises(ValueError, match="already contains students"):
        import_student_workbook(str(workbook_path), database_url)
