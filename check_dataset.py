#!/usr/bin/env python3
import openpyxl
import os

file_path = "/Users/diyanandkoomar/Downloads/NeuroSync_University_Intelligence_Dataset 2.xlsx"

if os.path.exists(file_path):
    print(f"✓ File found")
    print(f"  File size: {os.path.getsize(file_path) / 1024 / 1024:.2f} MB\n")
    
    wb = openpyxl.load_workbook(file_path)
    print(f"Sheets: {wb.sheetnames}\n")
    
    ws = wb.active
    if ws is None:
        print("No active worksheet found")
        raise SystemExit(1)
    print(f"Active sheet: {ws.title}")
    print(f"Rows: {ws.max_row}, Columns: {ws.max_column}")
    
    # Get headers
    headers = [cell.value for cell in ws[1]]
    print(f"\nHeaders ({len(headers)} total):")
    for i, h in enumerate(headers, 1):
        print(f"  {i}. {h}")
    
    # Get first data row
    if ws.max_row > 1:
        print(f"\nFirst student row:")
        for i, h in enumerate(headers, 1):
            cell = ws.cell(2, i)
            print(f"  {h}: {cell.value}")
else:
    print(f"✗ File not found at {file_path}")
